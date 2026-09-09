import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document

from job_hunter.cv_parser import (
    CVPage,
    detect_cv_signals,
    extract_cv_text,
    extract_cv_text_from_docx_bytes,
    extract_cv_text_from_legacy_doc_bytes,
    extract_text_from_pages,
)
from job_hunter.cv_store import CVStore


class CVParserTest(unittest.TestCase):
    def test_extract_text_from_pages_joins_non_empty_pdf_page_text(self):
        text = extract_text_from_pages(
            [
                CVPage("Muhamad Agustiar Falahi"),
                CVPage("Power BI, SSRS, BigQuery"),
                CVPage(None),
            ]
        )

        self.assertEqual(text, "Muhamad Agustiar Falahi\n\nPower BI, SSRS, BigQuery")

    def test_detect_cv_signals_returns_primary_and_bonus_matches(self):
        preferences = {
            "primary_keywords": ["Power BI", "SSRS", "BigQuery"],
            "bonus_keywords": ["Python", "Docker", "Airflow"],
        }

        signals = detect_cv_signals("Built Power BI dashboards with Python scripts.", preferences)

        self.assertEqual(signals.primary_matches, ("Power BI",))
        self.assertEqual(signals.bonus_matches, ("Python",))
        self.assertEqual(signals.total_matches, 2)

    def test_cv_store_saves_and_loads_extracted_text(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            store = CVStore(Path(tmpdir))

            store.save_text("Power BI and BigQuery")

            self.assertEqual(store.load_text(), "Power BI and BigQuery")

    def test_extract_cv_text_from_docx_bytes_reads_paragraphs_and_tables(self):
        doc = Document()
        doc.add_paragraph("Power BI Developer")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "SSRS"
        table.cell(0, 1).text = "BigQuery"
        stream = BytesIO()
        doc.save(stream)

        text = extract_cv_text_from_docx_bytes(stream.getvalue())

        self.assertIn("Power BI Developer", text)
        self.assertIn("SSRS", text)
        self.assertIn("BigQuery", text)

    def test_extract_cv_text_routes_by_filename_extension(self):
        doc = Document()
        doc.add_paragraph("Python and Airflow")
        stream = BytesIO()
        doc.save(stream)

        text = extract_cv_text(stream.getvalue(), "resume.docx")

        self.assertIn("Python and Airflow", text)

    def test_extract_cv_text_from_legacy_doc_bytes_returns_readable_best_effort_text(self):
        content = b"\x00\x01Power BI\x00\x00SSRS\x00BigQuery\x00"

        text = extract_cv_text_from_legacy_doc_bytes(content)

        self.assertIn("Power BI", text)
        self.assertIn("SSRS", text)

    def test_extract_cv_text_rejects_unsupported_extension(self):
        with self.assertRaises(ValueError):
            extract_cv_text(b"plain text", "resume.txt")


if __name__ == "__main__":
    unittest.main()
