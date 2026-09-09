"""CV text extraction and signal detection."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class CVPage:
    text: str | None

    def extract_text(self) -> str | None:
        return self.text


@dataclass(frozen=True)
class CVSignals:
    primary_matches: tuple[str, ...]
    bonus_matches: tuple[str, ...]

    @property
    def total_matches(self) -> int:
        return len(self.primary_matches) + len(self.bonus_matches)


def extract_cv_text_from_pdf_bytes(content: bytes) -> str:
    if not content:
        return ""
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    return extract_text_from_pages(reader.pages)


def extract_cv_text_from_docx_bytes(content: bytes) -> str:
    if not content:
        return ""
    from docx import Document

    document = Document(BytesIO(content))
    chunks: list[str] = []
    chunks.extend(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                chunks.append(row_text)
    return "\n\n".join(chunks)


def extract_cv_text_from_legacy_doc_bytes(content: bytes) -> str:
    if not content:
        return ""
    decoded_options = []
    for encoding in ("utf-8", "utf-16le", "latin-1"):
        try:
            decoded_options.append(content.decode(encoding, errors="ignore"))
        except LookupError:
            continue
    best = max(decoded_options, key=_readable_score, default="")
    cleaned = "".join(char if char.isprintable() or char in "\n\t " else " " for char in best)
    return " ".join(cleaned.split())


def extract_cv_text(content: bytes, filename: str) -> str:
    extension = Path(filename).suffix.casefold()
    if extension == ".pdf":
        return extract_cv_text_from_pdf_bytes(content)
    if extension == ".docx":
        return extract_cv_text_from_docx_bytes(content)
    if extension == ".doc":
        return extract_cv_text_from_legacy_doc_bytes(content)
    raise ValueError(f"Unsupported CV file type: {extension or 'unknown'}")


def extract_text_from_pages(pages: Any) -> str:
    chunks = []
    for page in pages:
        text = page.extract_text()
        if text and text.strip():
            chunks.append(text.strip())
    return "\n\n".join(chunks)


def detect_cv_signals(text: str, preferences: dict[str, object]) -> CVSignals:
    return CVSignals(
        primary_matches=_matches(text, preferences.get("primary_keywords", ())),
        bonus_matches=_matches(text, preferences.get("bonus_keywords", ())),
    )


def _matches(text: str, keywords: object) -> tuple[str, ...]:
    normalized = text.casefold()
    return tuple(str(keyword) for keyword in keywords or () if str(keyword).casefold() in normalized)


def _readable_score(text: str) -> int:
    return sum(1 for char in text if char.isalpha() or char.isspace())
